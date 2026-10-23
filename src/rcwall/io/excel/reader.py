"""Read Spanish worksheets, assemble canonical records and reuse canonical validators."""

import re
from collections import defaultdict
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from pydantic import ValidationError

from rcwall.domain.models import ProjectDefinition
from rcwall.validation.project import (
    validate_definition,
    validate_geometry,
    validate_references,
)

from .diagnostics import (
    CellLocation,
    ExcelImportResult,
    ExcelIssue,
    ExcelValidationError,
    ExcelValidationReport,
    SourceMap,
)
from .mapping import FIRST_DATA_ROW, HEADER_ROW, MAX_NATIVE_INTEGER, SHEETS, get_path, set_path


def _scan(workbook, issues):
    tables = {}
    for spec in SHEETS:
        rows = []
        tables[spec.collection] = rows
        if spec.name not in workbook.sheetnames:
            issues.append(
                ExcelIssue(
                    CellLocation(spec.name, 1, "Hoja"),
                    "missing_sheet",
                    "Falta la hoja obligatoria.",
                )
            )
            continue
        ws = workbook[spec.name]
        known = {c.header: c for c in spec.columns}
        positions = {}
        for cell in ws[HEADER_ROW]:
            if cell.value is None:
                continue
            if cell.value not in known:
                issues.append(
                    ExcelIssue(
                        CellLocation(spec.name, HEADER_ROW, cell.column_letter),
                        "unknown_header",
                        "Encabezado desconocido; conserve los encabezados de la plantilla.",
                    )
                )
            elif cell.value in positions:
                issues.append(
                    ExcelIssue(
                        CellLocation(spec.name, HEADER_ROW, str(cell.value)),
                        "duplicate_header",
                        "Encabezado repetido.",
                    )
                )
            else:
                positions[cell.value] = cell.column
        for column in spec.columns:
            if column.header not in positions:
                issues.append(
                    ExcelIssue(
                        CellLocation(spec.name, HEADER_ROW, column.header),
                        "missing_header",
                        "Falta el encabezado de esta columna.",
                    )
                )
        for row in ws.iter_rows(min_row=FIRST_DATA_ROW):
            if all(cell.value is None for cell in row):
                continue
            data, locations = {}, {}
            for column in spec.columns:
                location = CellLocation(spec.name, row[0].row, column.header)
                locations[column.field] = location
                index = positions.get(column.header)
                value = row[index - 1].value if index else None
                if value is None:
                    continue
                cell = row[index - 1]
                code = message = None
                if cell.data_type in ("f", "e"):
                    code, message = (
                        "formula_not_allowed",
                        "No se admiten fórmulas ni errores de Excel como datos.",
                    )
                elif column.kind == "decimal" and not isinstance(value, str):
                    code, message = (
                        "decimal_must_be_text",
                        "Guarde el decimal como texto; un número nativo de Excel "
                        "no garantiza conservar la precisión declarada.",
                    )
                elif column.kind == "integer":
                    if isinstance(value, str) and re.fullmatch(r"[+-]?[0-9]+", value):
                        value = int(value)
                    elif type(value) is not int or abs(value) > MAX_NATIVE_INTEGER:
                        code, message = (
                            "unsafe_integer",
                            "Escriba un entero; los valores con más de 15 dígitos "
                            "deben almacenarse como texto.",
                        )
                elif column.kind == "text" and not isinstance(value, str):
                    code, message = (
                        "text_required",
                        "Se requiere texto; no convierta identificadores o etiquetas en números.",
                    )
                if not code and column.choices:
                    reverse = {v: k for k, v in column.choices.items()}
                    if value not in reverse:
                        code, message = (
                            "invalid_option",
                            "Seleccione una opción en español de la lista desplegable.",
                        )
                    else:
                        value = reverse[value]
                if code:
                    issues.append(ExcelIssue(location, code, message))
                    continue
                set_path(data, column.path, value)
            for cell in row:
                if cell.value is not None and cell.column not in positions.values():
                    issues.append(
                        ExcelIssue(
                            CellLocation(spec.name, cell.row, cell.column_letter),
                            "unmapped_cell",
                            "Hay datos en una columna sin encabezado reconocido.",
                        )
                    )
            rows.append((data, locations))
    for name in workbook.sheetnames:
        if name not in {s.name for s in SHEETS}:
            issues.append(
                ExcelIssue(
                    CellLocation(name, 1, "Hoja"),
                    "unexpected_sheet",
                    "Hoja no admitida por esta interfaz de 17 hojas.",
                )
            )
    return tables


def _assemble(tables, issues, source):
    data = {}
    parent_index = defaultdict(list)
    for spec in SHEETS:
        if spec.model is None and spec.collection != "project":
            continue
        records = tables[spec.collection]
        if spec.collection == "project":
            if len(records) != 1:
                issues.append(
                    ExcelIssue(
                        CellLocation("Proyecto", FIRST_DATA_ROW, "ID de proyecto"),
                        "project_row_count",
                        "Diligencie exactamente una fila de proyecto.",
                    )
                )
            if records:
                data.update(records[0][0])
        else:
            data[spec.collection] = [row[0] for row in records]
        for index, (record, locations) in enumerate(records):
            prefix = () if spec.collection == "project" else (spec.collection, index)
            source.cells[prefix] = locations[spec.columns[0].field]
            for column in spec.columns:
                path = prefix + column.path
                source.cells[path] = locations[column.field]
                for length in range(len(prefix) + 1, len(path)):
                    source.cells.setdefault(path[:length], locations[column.field])
            if spec.collection != "project":
                identifier = get_path(record, spec.columns[0].path)
                if identifier is not None:
                    parent_index[(spec.collection, identifier)].append((record, prefix))
            if (
                spec.collection in ("geometries", "regions")
                and get_path(record, ("shape", "kind")) == "POLYGON"
            ):
                record["shape"].update(outer=[], holes=[])
            if (
                spec.collection == "bar_layouts"
                and get_path(record, ("distribution", "pattern")) == "EXPLICIT"
            ):
                record["distribution"]["bars"] = []

    def parent(collection, identifier, locations, field):
        candidates = parent_index.get((collection, identifier), [])
        if len(candidates) != 1:
            issues.append(
                ExcelIssue(
                    locations[field],
                    "advanced_parent",
                    "El propietario de esta fila no existe o su identificador está duplicado.",
                )
            )
            return None
        return candidates[0]

    groups = defaultdict(list)
    for collection, id_field, target in (
        ("explicit_bars", "layout_id", "bar_layouts"),
        ("material_properties", "section_id", "sections"),
        ("assembly_bars", "assembly_id", "transverse_assemblies"),
        ("vertices", "owner_id", None),
    ):
        for record, locations in tables[collection]:
            spec = next(s for s in SHEETS if s.collection == collection)
            missing = False
            for column in spec.columns:
                if get_path(record, column.path) is None:
                    issues.append(
                        ExcelIssue(
                            locations[column.field],
                            "missing",
                            "Campo obligatorio vacío en la fila avanzada.",
                        )
                    )
                    missing = True
            if missing:
                continue
            actual_target = record.get("owner_type") if collection == "vertices" else target
            owner = parent(actual_target, record.get(id_field), locations, id_field)
            if owner is None:
                continue
            owner_data, prefix = owner
            if collection == "vertices":
                if get_path(owner_data, ("shape", "kind")) != "POLYGON":
                    issues.append(
                        ExcelIssue(
                            locations[id_field],
                            "incompatible_children",
                            "Solo una geometría poligonal admite vértices.",
                        )
                    )
                    continue
                role, loop_order = record["loop_role"], record["loop_order"]
                if loop_order < 1 or (role == "outer" and loop_order != 1):
                    issues.append(
                        ExcelIssue(
                            locations["loop_order"],
                            "invalid_order",
                            "Exterior: orden 1. Huecos: enteros consecutivos desde 1.",
                        )
                    )
                    continue
                group = (collection, prefix, role, loop_order)
            elif collection == "explicit_bars":
                if get_path(owner_data, ("distribution", "pattern")) != "EXPLICIT":
                    issues.append(
                        ExcelIssue(
                            locations[id_field],
                            "incompatible_children",
                            "Las barras individuales requieren distribución explícita.",
                        )
                    )
                    continue
                group = (collection, prefix)
            elif collection == "material_properties":
                group = (collection, prefix, record["material_role"])
            else:
                group = (collection, prefix)
            groups[group].append((record, locations, owner_data))

    hole_orders = defaultdict(set)
    for group, entries in groups.items():
        entries.sort(key=lambda item: item[0]["order"])
        for expected, (record, locations, _) in enumerate(entries, 1):
            if record["order"] != expected:
                issues.append(
                    ExcelIssue(
                        locations["order"],
                        "invalid_order",
                        "El orden debe ser consecutivo desde 1, sin duplicados ni saltos.",
                    )
                )
        collection, prefix = group[:2]
        owner_data = entries[0][2]
        if collection == "vertices":
            role, loop_order = group[2:]
            if role == "outer":
                destination = owner_data["shape"]["outer"]
                path = prefix + ("shape", "outer")
            else:
                hole_orders[prefix].add(loop_order)
                holes = owner_data["shape"]["holes"]
                # Allocate by supplied contour count, never by a potentially huge input index.
                ranks = sorted({g[3] for g in groups if g[:3] == (collection, prefix, "holes")})
                rank = ranks.index(loop_order)
                while len(holes) <= rank:
                    holes.append([])
                destination = holes[rank]
                path = prefix + ("shape", "holes", rank)
            fields = {"x_mm": ("x_mm",), "y_mm": ("y_mm",)}
        elif collection == "explicit_bars":
            destination = owner_data["distribution"]["bars"]
            path = prefix + ("distribution", "bars")
            fields = {
                "bar_id": ("bar_id",),
                "position.x_mm": ("position", "x_mm"),
                "position.y_mm": ("position", "y_mm"),
                "diameter_mm": ("diameter_mm",),
            }
        elif collection == "material_properties":
            role = group[2]
            destination = (
                owner_data.setdefault("materials", {})
                .setdefault(role, {})
                .setdefault("properties", [])
            )
            path = prefix + ("materials", role, "properties")
            fields = {"name": ("name",), "value": ("value",), "unit": ("unit",)}
        else:
            destination = owner_data.setdefault("bent_bar_ids", [])
            path = prefix + ("bent_bar_ids",)
            fields = {}
        source.cells[path] = entries[0][1]["order"]
        for i, (record, locations, _) in enumerate(entries):
            source.cells[path + (i,)] = locations["order"]
            if collection == "assembly_bars":
                destination.append(record["bent_bar_id"])
                source.cells[path + (i,)] = locations["bent_bar_id"]
            else:
                item = {}
                for field, suffix in fields.items():
                    set_path(item, suffix, get_path(record, tuple(field.split("."))))
                    source.cells[path + (i,) + suffix] = locations[field]
                destination.append(item)
    for prefix, orders in hole_orders.items():
        if sorted(orders) != list(range(1, len(orders) + 1)):
            issues.append(
                ExcelIssue(
                    source.locate(prefix + ("shape", "holes", 0)),
                    "invalid_hole_order",
                    "Numere los huecos consecutivamente desde 1.",
                )
            )
    return data


def _validate(data, source, issues):
    outcome = validate_definition(data)
    issues.extend(source.translate(issue) for issue in outcome.report.issues)
    if any(i.stage == "types" for i in outcome.report.issues):
        # Reuse validators on independently valid records so one type error does not hide
        # unrelated reference/topology errors. Never return this diagnostic-only partial model.
        clean = {key: data.get(key) for key in ("project_id", "schema_version", "units")}
        remap = {}
        for spec in SHEETS:
            if spec.model is None:
                continue
            clean[spec.collection] = []
            for old_index, record in enumerate(data.get(spec.collection, [])):
                try:
                    model = spec.model.model_validate(record)
                except ValidationError:
                    continue
                index = len(clean[spec.collection])
                clean[spec.collection].append(model)
                remap[(spec.collection, index)] = (spec.collection, old_index)
        try:
            partial = ProjectDefinition.model_validate(clean)
        except ValidationError:
            partial = None
        if partial:
            for issue in validate_references(partial) + validate_geometry(partial):
                path = remap.get(issue.path[:2], issue.path[:2]) + issue.path[2:]
                issues.append(source.translate(issue.model_copy(update={"path": path})))
        else:
            # Invalid project metadata must not conceal topology errors in valid geometries.
            partial = ProjectDefinition.model_construct(**clean)
            for issue in validate_geometry(partial):
                path = remap.get(issue.path[:2], issue.path[:2]) + issue.path[2:]
                issues.append(source.translate(issue.model_copy(update={"path": path})))
    report = ExcelValidationReport(tuple(dict.fromkeys(issues)))
    return ExcelImportResult(outcome.project if report.valid else None, report)


def inspect_workbook(path: str | Path) -> ExcelImportResult:
    """Return accepted project or all available localized adapter/canonical diagnostics."""
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        return ExcelImportResult(
            None,
            ExcelValidationReport(
                (
                    ExcelIssue(
                        CellLocation("Libro", 1, "Archivo"),
                        "invalid_extension",
                        "Se requiere un archivo .xlsx.",
                    ),
                )
            ),
        )
    try:
        workbook = load_workbook(path, data_only=False, keep_links=False)
    except (OSError, BadZipFile, ValueError, KeyError, ParseError):
        return ExcelImportResult(
            None,
            ExcelValidationReport(
                (
                    ExcelIssue(
                        CellLocation("Libro", 1, "Archivo"),
                        "unreadable_workbook",
                        "No se pudo abrir un libro XLSX válido.",
                    ),
                )
            ),
        )
    issues, source = [], SourceMap()
    try:
        tables = _scan(workbook, issues)
        data = _assemble(tables, issues, source)
        return _validate(data, source, issues)
    finally:
        workbook.close()


def read_project(path: str | Path) -> ProjectDefinition:
    """Read an XLSX; raise ExcelValidationError with Spanish cell diagnostics on failure."""
    result = inspect_workbook(path)
    if result.project is None:
        raise ExcelValidationError(result.report)
    return result.project


def validate_workbook(path: str | Path) -> ExcelValidationReport:
    """Validate the workbook without writing files or repairing its contents."""
    return inspect_workbook(path).report
