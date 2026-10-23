"""Workbook generation and projection from the existing canonical contract."""

from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from rcwall.domain.models import ProjectDefinition
from rcwall.io.json import read_project as read_json_project
from rcwall.validation import validate_definition

from .diagnostics import ExcelValidationError, ExcelValidationReport, SourceMap
from .mapping import FIRST_DATA_ROW, HEADER_ROW, MAX_NATIVE_INTEGER, SHEETS, get_path

COLORS = {
    "required": ("24476A", "EBF2F8"),
    "optional": ("526374", "F4F6F8"),
    "conditional": ("69538A", "F4F0F8"),
}


def _project_rows(project: ProjectDefinition | None) -> dict[str, list[dict]]:
    rows = {sheet.collection: [] for sheet in SHEETS}
    if project is None:
        return rows
    data = project.model_dump(mode="json")
    rows["project"] = [data]
    for sheet in SHEETS:
        if sheet.model is not None:
            rows[sheet.collection] = data[sheet.collection]
    for collection, key in (("geometries", "geometry_id"), ("regions", "region_id")):
        for entity in data[collection]:
            shape = entity["shape"]
            if shape["kind"] != "POLYGON":
                continue
            loops = [("outer", 1, shape["outer"])] + [
                ("holes", index, hole) for index, hole in enumerate(shape["holes"], 1)
            ]
            for role, loop_order, vertices in loops:
                for order, vertex in enumerate(vertices, 1):
                    rows["vertices"].append(
                        {
                            "owner_type": collection,
                            "owner_id": entity[key],
                            "loop_role": role,
                            "loop_order": loop_order,
                            "order": order,
                            **vertex,
                        }
                    )
    for layout in data["bar_layouts"]:
        if layout["distribution"]["pattern"] == "EXPLICIT":
            for order, bar in enumerate(layout["distribution"]["bars"], 1):
                rows["explicit_bars"].append(
                    {"layout_id": layout["layout_id"], "order": order, **bar}
                )
    for section in data["sections"]:
        for role, material in section["materials"].items():
            for order, prop in enumerate(material["properties"], 1):
                rows["material_properties"].append(
                    {
                        "section_id": section["section_id"],
                        "material_role": role,
                        "order": order,
                        **prop,
                    }
                )
    for assembly in data["transverse_assemblies"]:
        for order, identifier in enumerate(assembly["bent_bar_ids"], 1):
            rows["assembly_bars"].append(
                {"assembly_id": assembly["assembly_id"], "order": order, "bent_bar_id": identifier}
            )
    return rows


def _write_value(cell, value, kind):
    if value is None:
        return
    if kind == "integer" and abs(value) <= MAX_NATIVE_INTEGER:
        cell.value = value
    else:
        text = str(value)
        if any(0xD800 <= ord(char) <= 0xDFFF or ord(char) in (0xFFFE, 0xFFFF) for char in text):
            raise ValueError(
                "El texto contiene caracteres Unicode no admitidos por el formato XML."
            )
        # Excel otherwise truncates long strings or treats a leading '=' as a formula.
        if len(text.encode("utf-16-le")) // 2 > 32767:
            raise ValueError("El texto supera el límite de 32767 caracteres de una celda Excel.")
        if any(ord(char) < 32 and char not in "\t\n\r" for char in text):
            raise ValueError("El texto contiene caracteres de control no admitidos por Excel.")
        cell.value = text
        cell.data_type = "s"


def _build(project: ProjectDefinition | None, editable_rows: int) -> Workbook:
    if isinstance(editable_rows, bool) or not isinstance(editable_rows, int) or editable_rows < 1:
        raise ValueError("La cantidad de filas editables debe ser un entero positivo.")
    rows = _project_rows(project)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Interfaz de entrada de muros de concreto reforzado V0.1"
    workbook.properties.subject = "Definiciones canónicas; sin cálculo de refuerzo"
    workbook.properties.creator = "StructureLab"
    for spec in SHEETS:
        ws = workbook.create_sheet(spec.name)
        last_column = get_column_letter(len(spec.columns))
        last_row = FIRST_DATA_ROW + max(editable_rows, len(rows[spec.collection])) - 1
        if last_row > 1_048_576:
            raise ValueError("La cantidad de filas supera la capacidad de una hoja Excel.")
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = f"B{FIRST_DATA_ROW}"
        ws.sheet_properties.tabColor = "69538A" if spec.advanced else "24476A"
        lines = [
            f"{spec.name} · {'Hoja avanzada' if spec.advanced else 'Hoja principal'} · V0.1",
            spec.instructions,
            "Decimales como texto con punto decimal. Conserve todos los dígitos; "
            "no use números nativos de Excel.",
            "Edite desde la fila 7. Una fila vacía se omite. "
            "Puede desproteger sin contraseña para ampliar filas. "
            "La validación de datos no demuestra cumplimiento normativo.",
            "Azul: obligatorio. Gris: opcional. Violeta: depende de la variante. "
            "Las filas informativas están protegidas.",
        ]
        for row, text in enumerate(lines, 1):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(spec.columns))
            cell = ws.cell(row, 1, text)
            cell.font = Font(
                name="Arial", size=14 if row == 1 else 10, color="24384C", bold=row == 1
            )
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[row].height = 30 if row == 1 else 62 if row == 2 else 45
        for index, column in enumerate(spec.columns, 1):
            letter = get_column_letter(index)
            width = max(23, min(40, len(column.header) + 3))
            if column.field == "description" or "specification" in column.field:
                width = 48
            ws.column_dimensions[letter].width = width
            cell = ws.cell(HEADER_ROW, index, column.header)
            dark, light = COLORS[column.requirement]
            cell.fill = PatternFill("solid", fgColor=dark)
            cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            label = {
                "required": "Obligatorio",
                "optional": "Opcional",
                "conditional": "Condicional",
            }[column.requirement]
            help_text = f"{label}. " + (column.help or "Consulte las indicaciones de esta hoja.")
            if column.kind == "decimal":
                help_text += (
                    " Escriba texto decimal, por ejemplo 12.50. "
                    "No se redondea ni se convierte desde números de Excel."
                )
            if column.kind == "integer":
                help_text += " Escriba un entero. Si supera 15 dígitos, guárdelo como texto."
            cell.comment = Comment(help_text, "StructureLab")
            for row in range(FIRST_DATA_ROW, last_row + 1):
                cell = ws.cell(row, index)
                cell.protection = Protection(locked=False)
                cell.fill = PatternFill("solid", fgColor=light)
                cell.font = Font(name="Arial", size=10, color="24384C")
                cell.number_format = "@"
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                    horizontal="right" if column.kind != "text" else "left",
                )
            if column.choices:
                options = ",".join(column.choices.values())
                validation = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
                validation.errorTitle = "Valor no permitido"
                validation.error = "Seleccione una de las opciones de la lista."
                validation.promptTitle = "Opciones admitidas"
                validation.prompt = "Use la lista desplegable. Los valores se muestran en español."
                validation.showErrorMessage = validation.showInputMessage = True
                validation.errorStyle = "stop"
                ws.add_data_validation(validation)
                validation.add(f"{letter}{FIRST_DATA_ROW}:{letter}{last_row}")
        ws.row_dimensions[HEADER_ROW].height = 42
        for row in range(FIRST_DATA_ROW, last_row + 1):
            ws.row_dimensions[row].height = 42
        for row_index, data in enumerate(rows[spec.collection], FIRST_DATA_ROW):
            for index, column in enumerate(spec.columns, 1):
                value = get_path(data, column.path)
                if column.choices and value is not None:
                    value = column.choices[value]
                _write_value(ws.cell(row_index, index), value, column.kind)
        ws.auto_filter.ref = f"A{HEADER_ROW}:{last_column}{last_row}"
        ws.protection.sheet = True
        ws.protection.autoFilter = False
        ws.protection.sort = False
        ws.protection.insertRows = False
        ws.protection.selectLockedCells = False
        ws.protection.selectUnlockedCells = False
        ws.print_title_rows = f"1:{HEADER_ROW}"
        ws.print_options.horizontalCentered = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        last_print_row = max(FIRST_DATA_ROW, FIRST_DATA_ROW + len(rows[spec.collection]) - 1)
        ws.print_area = f"A1:{last_column}{last_print_row}"
    return workbook


def _save(workbook: Workbook, path: str | Path) -> None:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("El archivo de destino debe tener extensión .xlsx.")
    try:
        raw = BytesIO()
        workbook.save(raw)
        encoded = BytesIO()
        # Some Windows XML backends translate LF to CRLF; XML parsers also normalize CR.
        # Restore inline text from original cells before encoding XML character references.
        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        worksheets = {
            f"xl/worksheets/sheet{i}.xml": sheet for i, sheet in enumerate(workbook.worksheets, 1)
        }
        with ZipFile(raw) as original, ZipFile(encoded, "w", ZIP_DEFLATED) as target:
            for entry in original.infolist():
                payload = original.read(entry.filename)
                if entry.filename in worksheets:
                    root = ElementTree.fromstring(payload)
                    sheet = worksheets[entry.filename]
                    for element in root.iter(namespace + "c"):
                        if element.get("t") == "inlineStr":
                            text = element.find(f"{namespace}is/{namespace}t")
                            if text is not None:
                                text.text = sheet[element.get("r")].value
                    payload = ElementTree.tostring(root, encoding="utf-8")
                    payload = payload.replace(b"\r", b"&#13;")
                target.writestr(entry, payload)
        path.write_bytes(encoded.getvalue())
    finally:
        workbook.close()


def create_template(path: str | Path, *, editable_rows: int = 200) -> None:
    """Create an empty 17-sheet workbook; a project must be entered before import."""
    _save(_build(None, editable_rows), path)


def export_project(
    project: ProjectDefinition, path: str | Path, *, editable_rows: int = 200
) -> None:
    """Validate and export every canonical field, preserving decimals and collection order."""
    outcome = validate_definition(project)
    if outcome.project is None:
        source = SourceMap()
        report = ExcelValidationReport(tuple(source.translate(i) for i in outcome.report.issues))
        raise ExcelValidationError(report)
    # Build fully before saving, including Excel capacity/character checks.
    _save(_build(outcome.project, editable_rows), path)


def generate_example(json_path: str | Path, xlsx_path: str | Path) -> None:
    """Generate the filled workbook from an existing canonical JSON input."""
    export_project(read_json_project(json_path), xlsx_path)
