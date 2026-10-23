from copy import deepcopy

import pytest
from openpyxl import load_workbook

from rcwall.domain.models import ProjectDefinition
from rcwall.io.excel import (
    ExcelValidationError,
    create_template,
    export_project,
    inspect_workbook,
    read_project,
    validate_workbook,
)
from rcwall.io.excel.mapping import BY_NAME, FIRST_DATA_ROW, HEADER_ROW, SHEETS
from rcwall.io.json import dumps_project


def edit(path, sheet, field, value, row=FIRST_DATA_ROW):
    workbook = load_workbook(path)
    spec = BY_NAME[sheet]
    index = next(i for i, column in enumerate(spec.columns, 1) if column.field == field)
    workbook[sheet].cell(row, index).value = value
    workbook.save(path)
    workbook.close()


@pytest.fixture
def excel_path(project_data, tmp_path):
    path = tmp_path / "entrada.xlsx"
    export_project(ProjectDefinition.model_validate(project_data), path, editable_rows=5)
    return path


def test_empty_template_and_visible_spanish(tmp_path):
    path = tmp_path / "plantilla.xlsx"
    create_template(path, editable_rows=3)
    wb = load_workbook(path)
    assert len(wb.sheetnames) == 17
    assert wb.sheetnames == [s.name for s in SHEETS]
    assert "Trayectorias" not in wb.sheetnames
    for spec in SHEETS:
        ws = wb[spec.name]
        assert ws.sheet_state == "visible"
        assert [cell.value for cell in ws[HEADER_ROW]] == [c.header for c in spec.columns]
        assert all(cell.value is None for row in ws.iter_rows(min_row=7) for cell in row)
        assert ws.freeze_panes == "B7"
        assert ws.auto_filter.ref
        assert ws.protection.sheet
        assert ws["A6"].protection.locked
        assert not ws["A7"].protection.locked
        assert (
            "Hoja avanzada" in ws["A1"].value
            if spec.advanced
            else "Hoja principal" in ws["A1"].value
        )
        for index, _col in enumerate(spec.columns, 1):
            assert ws.cell(6, index).comment
            assert ws.cell(7, index).number_format == "@"
        for validation in ws.data_validations.dataValidation:
            assert validation.errorTitle == "Valor no permitido"
            assert "UNRESOLVED" not in validation.formula1
            assert "RECTANGLE" not in validation.formula1
    wb.close()
    report = validate_workbook(path)
    assert not report.valid  # Structurally valid blank template is not a filled project.
    assert "project_row_count" in {i.code for i in report.issues}


def test_example_roundtrip_and_decimal_storage(excel_path, project_data):
    original = ProjectDefinition.model_validate(project_data)
    assert dumps_project(read_project(excel_path)) == dumps_project(original)
    wb = load_workbook(excel_path)
    cell = wb["Geometrías"]["E7"]
    assert cell.value == "3000.000000000000000000000000001"
    assert cell.data_type == "s"
    assert wb["Barras"]["F7"].value == 6
    wb.close()


def rich_data(data):
    def points(values):
        return [{"x_mm": str(x), "y_mm": str(y)} for x, y in values]

    polygon = {
        "kind": "POLYGON",
        "outer": points([(0, 0), (100, 0), (100, 100), (0, 100)]),
        "holes": [
            points([(60, 60), (60, 80), (80, 80), (80, 60)]),
            points([(10, 10), (10, 20), (20, 20), (20, 10)]),
        ],
    }
    data["geometries"].append({"geometry_id": "GP", "shape": polygon})
    data["regions"] = [{"region_id": "RP", "geometry_id": "GP", "shape": polygon, "role": "OTHER"}]
    data["configurations"].append({"config_id": "CP", "geometry_id": "GP"})
    spacing = deepcopy(data["bar_layouts"][0])
    spacing["layout_id"] = "LS"
    spacing["distribution"].pop("quantity")
    spacing["distribution"].update(
        pattern="LINEAR_SPACING",
        spacing_mm="1.23000",
        spacing_policy="  EXACT user text  ",
        alignment="  START  ",
    )
    data["bar_layouts"].append(spacing)
    data["bar_layouts"].append(
        {
            "layout_id": "LE",
            "config_id": "CP",
            "region_id": "RP",
            "distribution": {
                "pattern": "EXPLICIT",
                "bars": [
                    {
                        "bar_id": "B2",
                        "position": {"x_mm": "1.000", "y_mm": "2"},
                        "diameter_mm": "16.00",
                    },
                    {
                        "bar_id": "B1",
                        "position": {"x_mm": "3", "y_mm": "4.00"},
                        "diameter_mm": "12",
                    },
                ],
            },
        }
    )
    data["sections"][0]["materials"]["concrete"]["properties"] = [
        {"name": "valor2", "value": "1E-40", "unit": "unidad libre"},
        {"name": "valor1", "value": "-0.000", "unit": "MPa"},
    ]
    data["sections"][0]["description"] = "=texto libre, no una fórmula\nSegunda línea"
    data["bent_bars"] = [
        {"bent_bar_id": "BB", "definition_status": "UNRESOLVED", "diameter_mm": "8.00"}
    ]
    data["transverse_assemblies"] = [
        {
            "assembly_id": "T",
            "config_id": "C1",
            "definition_status": "UNRESOLVED",
            "bent_bar_ids": ["BB"],
        }
    ]
    data["vertical_placements"] = [
        {
            "placement_id": "V",
            "assignment_id": "A1",
            "assembly_id": "T",
            "definition_status": "UNRESOLVED",
            "z_start_mm": "-3.00",
            "z_end_mm": "1.0",
            "spacing_mm": "0.100",
            "spacing_policy": "  libre  ",
        }
    ]
    data["drawings"] = [
        {
            "drawing_id": "D",
            "section_id": "S1",
            "definition_status": "UNRESOLVED",
            "template_id": "TEXTO",
            "drawing_scale": "25.000",
        }
    ]
    return data


def test_all_fields_order_and_variants_roundtrip(project_data, tmp_path):
    project = ProjectDefinition.model_validate(rich_data(project_data))
    path = tmp_path / "completo.xlsx"
    export_project(project, path, editable_rows=1)
    assert dumps_project(read_project(path)) == dumps_project(project)
    wb = load_workbook(path)
    assert wb["Secciones"]["F7"].data_type == "s"
    # Physical order of child rows does not override the explicit original order.
    ws = wb["Barras explícitas"]
    a, b = [c.value for c in ws[7]], [c.value for c in ws[8]]
    for index, value in enumerate(b, 1):
        ws.cell(7, index, value)
    for index, value in enumerate(a, 1):
        ws.cell(8, index, value)
    wb.save(path)
    wb.close()
    assert dumps_project(read_project(path)) == dumps_project(project)


@pytest.mark.parametrize(
    ("sheet", "field", "value", "code"),
    [
        ("Geometrías", "shape.length_mm", 3000.5, "decimal_must_be_text"),
        ("Geometrías", "shape.length_mm", 3000, "decimal_must_be_text"),
        ("Geometrías", "shape.length_mm", "0", "greater_than"),
        ("Geometrías", "shape.length_mm", "NaN", "decimal_input"),
        ("Geometrías", "shape.length_mm", "3,5", "decimal_input"),
        ("Geometrías", "shape.kind", "RECTANGLE", "invalid_option"),
        ("Barras", "distribution.quantity", None, "missing"),
        ("Barras", "distribution.quantity", 1.25, "unsafe_integer"),
        ("Barras", "distribution.quantity", True, "unsafe_integer"),
        ("Barras", "distribution.quantity", 1234567890123456, "unsafe_integer"),
        ("Barras", "distribution.quantity", 0, "greater_than"),
        ("Barras", "distribution.spacing_mm", "10", "extra_forbidden"),
        ("Asignaciones", "section_id", "NO-EXISTE", "missing_reference"),
        ("Muros", "wall_id", "G1", "duplicate_id"),
        ("Muros", "wall_id", 100, "text_required"),
        ("Muros", "description", "=1+1", "formula_not_allowed"),
        ("Proyecto", "units", "mm", "invalid_option"),
        ("Proyecto", "schema_version", "9.0", "invalid_option"),
    ],
)
def test_localized_errors(excel_path, sheet, field, value, code):
    edit(excel_path, sheet, field, value)
    result = inspect_workbook(excel_path)
    assert result.project is None
    header = next(c.header for c in BY_NAME[sheet].columns if c.field == field)
    assert any(
        i.code == code
        and i.location.sheet == sheet
        and i.location.row == 7
        and i.location.column == header
        for i in result.report.issues
    )
    assert "Errores" in result.report.to_dict()


def test_blank_rows_preserve_physical_error_location(excel_path):
    wb = load_workbook(excel_path)
    wb["Barras"].insert_rows(7)
    wb["Barras"]["F8"] = None
    wb.save(excel_path)
    wb.close()
    result = inspect_workbook(excel_path)
    assert any(
        i.location.sheet == "Barras" and i.location.row == 8 and i.location.column == "Cantidad"
        for i in result.report.issues
    )


@pytest.mark.parametrize(
    "change",
    ["missing_sheet", "missing_header", "duplicate_header", "unknown_header", "unexpected_sheet"],
)
def test_layout_errors(excel_path, change):
    wb = load_workbook(excel_path)
    if change == "missing_sheet":
        del wb["Muros"]
    elif change == "missing_header":
        wb["Muros"]["B6"] = None
    elif change == "duplicate_header":
        wb["Muros"]["B6"] = "ID de muro"
    elif change == "unknown_header":
        wb["Muros"]["B6"] = "Campo desconocido"
    else:
        wb.create_sheet("Trayectorias")
    wb.save(excel_path)
    wb.close()
    assert change in {i.code for i in validate_workbook(excel_path).issues}


def test_large_quantity_text_roundtrip(project_data, tmp_path):
    project_data["bar_layouts"][0]["distribution"]["quantity"] = 1234567890123456789012345
    project = ProjectDefinition.model_validate(project_data)
    path = tmp_path / "grande.xlsx"
    export_project(project, path, editable_rows=1)
    assert dumps_project(read_project(path)) == dumps_project(project)
    wb = load_workbook(path)
    assert wb["Barras"]["F7"].data_type == "s"
    wb.close()


def test_multiple_independent_errors_survive(project_data, tmp_path):
    project = ProjectDefinition.model_validate(rich_data(project_data))
    path = tmp_path / "errores.xlsx"
    export_project(project, path, editable_rows=1)
    edit(path, "Barras", "distribution.quantity", None)
    edit(path, "Asignaciones", "section_id", "AUSENTE")
    edit(path, "Vértices", "x_mm", "100", row=7)
    codes = {i.code for i in validate_workbook(path).issues}
    assert {"missing", "missing_reference", "repeated_vertex"} <= codes


@pytest.mark.parametrize(
    ("sheet", "field", "value", "code"),
    [
        ("Vértices", "owner_id", "AUSENTE", "advanced_parent"),
        ("Vértices", "order", 2, "invalid_order"),
        ("Vértices", "owner_id", "G1", "incompatible_children"),
        ("Barras explícitas", "layout_id", "LAYOUT1", "incompatible_children"),
        ("Barras explícitas", "order", 0, "invalid_order"),
        ("Propiedades de materiales", "value", 1.1, "decimal_must_be_text"),
        ("Estribos-barras", "bent_bar_id", "AUSENTE", "missing_reference"),
        ("Distribución vertical", "z_end_mm", "-4", "invalid_interval"),
    ],
)
def test_advanced_diagnostics(project_data, tmp_path, sheet, field, value, code):
    path = tmp_path / "avanzado.xlsx"
    export_project(ProjectDefinition.model_validate(rich_data(project_data)), path, editable_rows=1)
    edit(path, sheet, field, value)
    assert any(i.code == code and i.location.sheet == sheet for i in validate_workbook(path).issues)


def test_empty_collections_optional_nulls_and_reused_configuration(project_data, excel_path):
    result = read_project(excel_path)
    assert result.bent_bars == ()
    assert result.geometries[0].description is None
    assert result.sections[0].config_id == result.sections[1].config_id
    assert dumps_project(result) == dumps_project(ProjectDefinition.model_validate(project_data))


def test_editing_valid_data_is_reflected(excel_path):
    edit(excel_path, "Muros", "description", " Muro editado por el usuario ")
    assert read_project(excel_path).walls[0].description == " Muro editado por el usuario "


def test_no_mutation_on_invalid_geometry(project_data, tmp_path):
    path = tmp_path / "geometria.xlsx"
    export_project(ProjectDefinition.model_validate(rich_data(project_data)), path, editable_rows=1)
    edit(path, "Vértices", "x_mm", "100")
    before = path.read_bytes()
    with pytest.raises(ExcelValidationError):
        read_project(path)
    assert path.read_bytes() == before


def test_invalid_project_does_not_overwrite_file(excel_path, project_data):
    before = excel_path.read_bytes()
    project_data["assignments"][0]["section_id"] = "AUSENTE"
    with pytest.raises(ExcelValidationError):
        export_project(ProjectDefinition.model_validate(project_data), excel_path)
    assert excel_path.read_bytes() == before


def test_unreadable_and_wrong_extension(tmp_path):
    path = tmp_path / "roto.xlsx"
    path.write_text("no es un archivo Excel", encoding="utf-8")
    assert validate_workbook(path).issues[0].code == "unreadable_workbook"
    assert validate_workbook(tmp_path / "entrada.xls").issues[0].code == "invalid_extension"


def test_all_canonical_leaf_fields_have_explicit_columns():
    from rcwall.io.json import project_json_schema

    schema = project_json_schema()

    def leaves(node, prefix=()):
        if "$ref" in node:
            yield from leaves(schema["$defs"][node["$ref"].split("/")[-1]], prefix)
        elif "properties" in node:
            for name, child in node["properties"].items():
                yield from leaves(child, prefix + (name,))
        elif "oneOf" in node:
            for child in node["oneOf"]:
                yield from leaves(child, prefix)
        elif node.get("type") != "array":
            yield ".".join(prefix)

    for sheet in SHEETS:
        if sheet.model is None:
            continue
        expected = set(leaves({"$ref": f"#/$defs/{sheet.model.__name__}"}))
        assert expected == {col.field for col in sheet.columns}


def test_translation_tables_are_bijective_and_match_canonical_choices():
    from rcwall.io.excel.mapping import MATERIALS, PATTERNS, ROLES, SHAPES, STATUS
    from rcwall.io.json import project_json_schema

    definitions = project_json_schema()["$defs"]
    assert set(SHAPES) == {
        definitions[n]["properties"]["kind"]["const"]
        for n in ("RectangleDefinition", "PolygonDefinition")
    }
    assert set(PATTERNS) == {
        definitions[n]["properties"]["pattern"]["const"]
        for n in ("LinearCount", "LinearSpacing", "ExplicitLayout")
    }
    assert set(ROLES) == set(definitions["RegionDefinition"]["properties"]["role"]["enum"])
    assert set(STATUS) == {"UNRESOLVED"}
    assert set(MATERIALS) == set(definitions["SectionMaterials"]["properties"])
    for sheet in SHEETS:
        for column in sheet.columns:
            if column.choices:
                assert len(set(column.choices.values())) == len(column.choices)


def test_reordered_columns_are_read_by_header(excel_path, project_data):
    wb = load_workbook(excel_path)
    ws = wb["Muros"]
    for row in range(6, ws.max_row + 1):
        left, right = ws.cell(row, 1).value, ws.cell(row, 2).value
        ws.cell(row, 1).value = right
        ws.cell(row, 2).value = left
    wb.save(excel_path)
    wb.close()
    assert dumps_project(read_project(excel_path)) == dumps_project(
        ProjectDefinition.model_validate(project_data)
    )


def test_free_text_unicode_and_line_endings_are_preserved(project_data, tmp_path):
    project_data["walls"][0]["description"] = "  Texto\r\nÁrbol\t漢字😀  "
    path = tmp_path / "texto.xlsx"
    project = ProjectDefinition.model_validate(project_data)
    export_project(project, path, editable_rows=1)
    assert dumps_project(read_project(path)) == dumps_project(project)


def test_no_macro_or_generated_formula(excel_path):
    from zipfile import ZipFile

    with ZipFile(excel_path) as archive:
        assert not any("vba" in name.lower() for name in archive.namelist())
    wb = load_workbook(excel_path)
    assert all(cell.data_type != "f" for ws in wb for row in ws for cell in row)
    wb.close()


def test_hole_order_gap_and_orphan_material_are_not_discarded(project_data, tmp_path):
    path = tmp_path / "orden.xlsx"
    export_project(ProjectDefinition.model_validate(rich_data(project_data)), path, editable_rows=1)
    wb = load_workbook(path)
    ws = wb["Vértices"]
    # First outer uses rows 7..10; first hole rows 11..14.
    for row in range(11, 15):
        ws.cell(row, 4).value = 3
    wb["Propiedades de materiales"]["A7"] = "HUÉRFANA"
    wb.save(path)
    wb.close()
    codes = {issue.code for issue in validate_workbook(path).issues}
    assert {"invalid_hole_order", "advanced_parent"} <= codes


@pytest.mark.parametrize("value", ["a" * 32768, "control\x01"], ids=["long_text", "control"])
def test_excel_text_capacity_errors_do_not_overwrite(project_data, excel_path, value):
    project_data["walls"][0]["description"] = value
    before = excel_path.read_bytes()
    with pytest.raises(ValueError):
        export_project(ProjectDefinition.model_validate(project_data), excel_path, editable_rows=1)
    assert excel_path.read_bytes() == before
